"""
OSC utility functions extracted from api/server.py.

Database config/connection, JSON/path/folder helpers, file reading,
skill execution, parsing/normalization, and core helper functions.
"""

from __future__ import annotations

import difflib
import errno
import functools
import json
import logging
import os
import platform
import re
import shutil
import subprocess
import sys
import threading
import urllib.error
import urllib.request
import html as ihtml
from contextlib import contextmanager
from datetime import date, datetime, time as dt_time, timedelta
from decimal import Decimal
from pathlib import Path
from typing import TYPE_CHECKING, Optional

try:
    from flask_login import current_user
except Exception:
    class _NoUser:
        is_authenticated = False
        username = None

    current_user = _NoUser()

from api.runtime_paths import get_config_path, get_magi_root_dir, get_orch_dir
from api.case_path_mapper import (
    local_synology_path_candidates,
    default_synology_share_roots,
    preferred_synology_share_roots,
    translate_local_path_to_canonical,
)
from api.osc.insight_filters import is_non_extractable_legal_insight

try:
    from api.tw_output_guard import normalize_output_text as _normalize_output_text
except Exception:
    _normalize_output_text = None

logger = logging.getLogger(__name__)


_OSC_FILE_STAGE_MAX_CONCURRENT = max(
    1,
    int(os.environ.get("PAPERCLIP_FILE_STAGE_MAX_CONCURRENT", "1") or "1"),
)
_OSC_FILE_STAGE_ACQUIRE_TIMEOUT = max(
    0.0,
    float(os.environ.get("PAPERCLIP_FILE_STAGE_ACQUIRE_TIMEOUT_SEC", "0.05") or "0.05"),
)
_OSC_FILE_STAGE_SEMAPHORE = threading.BoundedSemaphore(_OSC_FILE_STAGE_MAX_CONCURRENT)
_OSC_FILE_STAGE_LOCAL = threading.local()
_OSC_DIRECTORY_IO_MAX_CONCURRENT = max(
    1,
    # Gateway has eight bounded workers. Six concurrent read-only directory
    # operations preserve two workers for health/login/API traffic while
    # allowing several people to browse or preview files at the same time.
    int(os.environ.get("PAPERCLIP_DIRECTORY_IO_MAX_CONCURRENT", "6") or "6"),
)
_OSC_DIRECTORY_IO_ACQUIRE_TIMEOUT = max(
    0.0,
    float(os.environ.get("PAPERCLIP_DIRECTORY_IO_ACQUIRE_TIMEOUT_SEC", "0.75") or "0.75"),
)
_OSC_DIRECTORY_IO_SEMAPHORE = threading.BoundedSemaphore(_OSC_DIRECTORY_IO_MAX_CONCURRENT)
_OSC_DIRECTORY_IO_LOCAL = threading.local()


@contextmanager
def _osc_file_stage_slot():
    """Process-wide, re-entrant bulkhead for NAS-to-local file staging."""
    depth = int(getattr(_OSC_FILE_STAGE_LOCAL, "depth", 0) or 0)
    if depth > 0:
        _OSC_FILE_STAGE_LOCAL.depth = depth + 1
        try:
            yield
        finally:
            _OSC_FILE_STAGE_LOCAL.depth = depth
        return

    acquired = _OSC_FILE_STAGE_SEMAPHORE.acquire(timeout=_OSC_FILE_STAGE_ACQUIRE_TIMEOUT)
    if not acquired:
        raise OSError(errno.EBUSY, "file staging busy; retry shortly")
    _OSC_FILE_STAGE_LOCAL.depth = 1
    try:
        yield
    finally:
        _OSC_FILE_STAGE_LOCAL.depth = 0
        _OSC_FILE_STAGE_SEMAPHORE.release()


def _osc_stage_bulkhead(func):
    """Decorate a staging entrypoint without deadlocking nested fallbacks."""
    @functools.wraps(func)
    def wrapped(*args, **kwargs):
        with _osc_file_stage_slot():
            return func(*args, **kwargs)

    return wrapped


@contextmanager
def _osc_directory_io_slot():
    """Keep slow NAS directory operations from occupying every HTTP worker."""
    depth = int(getattr(_OSC_DIRECTORY_IO_LOCAL, "depth", 0) or 0)
    if depth > 0:
        _OSC_DIRECTORY_IO_LOCAL.depth = depth + 1
        try:
            yield
        finally:
            _OSC_DIRECTORY_IO_LOCAL.depth = depth
        return

    acquired = _OSC_DIRECTORY_IO_SEMAPHORE.acquire(timeout=_OSC_DIRECTORY_IO_ACQUIRE_TIMEOUT)
    if not acquired:
        raise OSError(errno.EBUSY, "directory I/O busy; retry shortly")
    _OSC_DIRECTORY_IO_LOCAL.depth = 1
    try:
        yield
    finally:
        _OSC_DIRECTORY_IO_LOCAL.depth = 0
        _OSC_DIRECTORY_IO_SEMAPHORE.release()


if TYPE_CHECKING:
    import mysql.connector.pooling  # pragma: no cover - typing only
    import mysql.connector  # pragma: no cover - typing only


def _load_mysql():
    """Lazy-load mysql.connector only when DB helpers are actually used."""
    try:
        import mysql.connector
        import mysql.connector.pooling
        return mysql.connector, mysql.connector.pooling
    except Exception as exc:
        raise RuntimeError("OSC DB features require mysql package") from exc


def _load_local_dotenv() -> None:
    try:
        from dotenv import load_dotenv as _load_dotenv
        from api.runtime_paths import get_magi_root_dir
        _load_dotenv(get_magi_root_dir() / ".env", override=False)
    except Exception:
        logger.debug("silent-catch dotenv load", exc_info=True)


_load_local_dotenv()

_OSC_NAS_HOME_USER = (
    os.environ.get("MAGI_NAS_HOME_USER")
    or os.environ.get("MAGI_NAS_USER")
    or "home"
).strip().strip("/\\") or "home"
_OSC_NAS_CLOSED_SHARE_NAME = (
    os.environ.get("MAGI_NAS_CLOSED_SHARE_NAME")
    or os.environ.get("MAGI_NAS_ARCHIVE_SHARE")
    or "lumi"
).strip().strip("/\\") or "lumi"


def _osc_closed_share_aliases() -> list[str]:
    aliases = [_OSC_NAS_CLOSED_SHARE_NAME, "lumi", "archive"]
    out: list[str] = []
    for item in aliases:
        text = str(item or "").strip().strip("/\\")
        if text and text not in out:
            out.append(text)
    return out


def _osc_canonical_active_share_windows() -> str:
    return (
        os.environ.get("MAGI_CANONICAL_ACTIVE_SHARE_PREFIX")
        or f"Z:\\{_OSC_NAS_HOME_USER}"
    ).replace("/", "\\").rstrip("\\")


def _osc_canonical_active_share_posix() -> str:
    return _osc_canonical_active_share_windows().replace("\\", "/")

# ---------------------------------------------------------------------------
# 1. Database config and connection functions
# ---------------------------------------------------------------------------

OSC_WEB_DB_CONFIG = {
    "host": os.environ.get("OSC_DB_HOST", "") or os.environ.get("DB_HOST", "127.0.0.1"),
    "port": int(os.environ.get("OSC_DB_PORT", "") or os.environ.get("DB_PORT", "3306")),
    "user": os.environ.get("OSC_DB_USER", "") or os.environ.get("DB_USER", ""),
    "password": os.environ.get("OSC_DB_PASSWORD", "") or os.environ.get("DB_PASSWORD", ""),
    "database": os.environ.get("OSC_DB_NAME", "law_firm_data"),
}


def _load_code_db_profile(profile_name: str = "Studio_VPN_Remote"):
    cfg_paths = [
        str(get_config_path("config.json")),
        os.path.join(str(get_orch_dir()), "config.json"),
    ]
    for p in cfg_paths:
        try:
            if not os.path.exists(p):
                continue
            with open(p, "r", encoding="utf-8") as f:
                obj = json.load(f) or {}
            for it in (obj.get("mariadb_profiles") or []):
                if not isinstance(it, dict):
                    continue
                if str(it.get("profile_name") or "").strip() != profile_name:
                    continue
                c = (it.get("config") or {}) if isinstance(it.get("config"), dict) else {}
                return {
                    "host": str(c.get("host") or os.environ.get("OSC_DB_HOST", "127.0.0.1")),
                    "port": int(c.get("port") or 3306),
                    "user": str(c.get("user") or os.environ.get("OSC_DB_USER", "python_user")),
                    "password": str(c.get("password") or os.environ.get("OSC_DB_PASSWORD", "")),
                    "database": str(c.get("database") or "law_firm_data"),
                }
        except Exception:
            continue
    return None


def _resolve_osc_web_db_config():
    # Explicit OSC_WEB_DB_* always wins.
    explicit = any((os.environ.get(k, "") or "").strip() for k in ["OSC_WEB_DB_HOST", "OSC_WEB_DB_PORT", "OSC_WEB_DB_USER", "OSC_WEB_DB_PASSWORD", "OSC_WEB_DB_NAME"])
    if explicit:
        return {
            "host": os.environ.get("OSC_WEB_DB_HOST") or os.environ.get("OSC_DB_HOST", "127.0.0.1"),
            "port": int((os.environ.get("OSC_WEB_DB_PORT") or "3306").strip()),
            "user": os.environ.get("OSC_WEB_DB_USER") or "python_user",
            "password": os.environ.get("OSC_WEB_DB_PASSWORD") or "",
            "database": os.environ.get("OSC_WEB_DB_NAME") or "law_firm_data",
        }

    # Default behavior: prefer CRUD-capable profile from code config.
    prefer_crud_user = (os.environ.get("OSC_WEB_PREFER_CRUD_USER", "1") or "1").strip().lower() in {"1", "true", "yes", "on"}
    if prefer_crud_user:
        prof = _load_code_db_profile("Studio_VPN_Remote")
        if prof:
            return prof

    # Fallback to existing MAGI env chain.
    return {
        "host": os.environ.get("OSC_DB_HOST") or os.environ.get("MAGI_REMOTE_DB_HOST") or "127.0.0.1",
        "port": int((os.environ.get("OSC_DB_PORT") or os.environ.get("MAGI_REMOTE_DB_PORT") or "3306").strip()),
        "user": os.environ.get("OSC_DB_USER") or os.environ.get("MAGI_REMOTE_DB_USER") or "python_user",
        "password": os.environ.get("OSC_DB_PASSWORD") or os.environ.get("MAGI_REMOTE_DB_PASSWORD") or "",
        "database": os.environ.get("OSC_DB_NAME") or os.environ.get("MAGI_REMOTE_DB_NAME") or "law_firm_data",
    }


OSC_WEB_DB_CONFIG = _resolve_osc_web_db_config()


def _osc_web_db_candidates():
    """Build ordered candidate list, respecting db_failover's active host."""
    base = dict(OSC_WEB_DB_CONFIG)

    # Consult failover module for the *current* active host so the candidate
    # list is not frozen to whatever was resolved at import time.
    try:
        from api.db_failover import get_failover_status
        status = get_failover_status()
        if status.get("failover_active"):
            # Failover is active → local DB should be tried first.
            local_host = (os.environ.get("MAGI_LOCAL_DB_HOST") or "127.0.0.1").strip()
            local_port = int((os.environ.get("MAGI_LOCAL_DB_PORT") or "3306").strip())
            local_cfg = {
                "host": local_host,
                "port": local_port,
                "user": (os.environ.get("MAGI_LOCAL_DB_USER") or base["user"]).strip(),
                "password": os.environ.get("MAGI_LOCAL_DB_PASSWORD") or base["password"],
                "database": (os.environ.get("MAGI_LOCAL_DB_NAME") or base["database"]).strip(),
            }
            # Local first, then remote as fallback.
            return [local_cfg, base]
    except Exception:
        pass

    # Normal path: remote (from config profile) first, local as fallback.
    cands = [base]
    local_host = (os.environ.get("MAGI_LOCAL_DB_HOST") or "127.0.0.1").strip()
    local_port = int((os.environ.get("MAGI_LOCAL_DB_PORT") or "3306").strip())
    local_user = (os.environ.get("MAGI_LOCAL_DB_USER") or base["user"]).strip()
    local_pass = os.environ.get("MAGI_LOCAL_DB_PASSWORD") or base["password"]
    local_name = (os.environ.get("MAGI_LOCAL_DB_NAME") or base["database"]).strip()
    if (local_host, local_port, local_name, local_user) != (base["host"], base["port"], base["database"], base["user"]):
        cands.append(
            {
                "host": local_host,
                "port": local_port,
                "user": local_user,
                "password": local_pass,
                "database": local_name,
            }
        )
    return cands


# ---------------------------------------------------------------------------
# Connection pool management
# ---------------------------------------------------------------------------
_pool = None          # MySQLConnectionPool instance (lazily created)
_pool_cfg = None      # The db config dict that the pool was created with
_pool_lock = __import__("threading").Lock()
_pool_failover_active = None   # Track failover state when pool was created
_pool_seq = 0                  # Monotonic counter to generate unique pool names

_logger = logging.getLogger(__name__)


def _current_failover_active():
    """Return db_failover's current failover_active flag, or None if unavailable."""
    try:
        from api.db_failover import get_failover_status
        return get_failover_status().get("failover_active")
    except Exception:
        return None


def _get_pool():
    """Return (pool, cfg) – create the pool lazily on first call.

    Automatically resets the pool when the failover state has changed since the
    pool was created (e.g. switched from remote to local or vice-versa).
    """
    global _pool, _pool_cfg, _pool_failover_active, _pool_seq

    # Fast path: pool exists AND failover state hasn't changed.
    if _pool is not None:
        current_fo = _current_failover_active()
        if current_fo == _pool_failover_active:
            return _pool, _pool_cfg
        # Failover state changed → discard stale pool.
        _logger.info("Failover state changed (%s → %s), resetting OSC pool",
                      _pool_failover_active, current_fo)
        _reset_pool()

    with _pool_lock:
        # Double-check after acquiring lock
        if _pool is not None:
            current_fo = _current_failover_active()
            if current_fo == _pool_failover_active:
                return _pool, _pool_cfg
            _pool = None
            _pool_cfg = None

        # Let failover module update env before we pick candidates
        try:
            from api.db_failover import probe_remote
            probe_remote()
        except Exception:
            pass

        fo_state = _current_failover_active()
        last_err = None
        for cfg in _osc_web_db_candidates():
            try:
                _pool_seq += 1
                mysql_connector, mysql_pooling = _load_mysql()
                pool = mysql_pooling.MySQLConnectionPool(
                    pool_name=f"osc_pool_{_pool_seq}",
                    pool_size=5,
                    pool_reset_session=True,
                    host=cfg["host"],
                    port=int(cfg["port"]),
                    user=cfg["user"],
                    password=cfg["password"],
                    database=cfg["database"],
                    autocommit=False,
                    charset="utf8mb4",
                    collation="utf8mb4_unicode_ci",
                    connection_timeout=5,
                )
                _pool = pool
                _pool_cfg = cfg
                _pool_failover_active = fo_state
                _logger.info("OSC connection pool created: host=%s port=%s db=%s (failover=%s)",
                             cfg["host"], cfg["port"], cfg["database"], fo_state)
                return _pool, _pool_cfg
            except Exception as e:
                last_err = e
                continue
        raise last_err or RuntimeError("osc_web_db_connect_failed")


def _reset_pool():
    """Discard the current pool so the next call to ``_get_pool`` rebuilds it."""
    global _pool, _pool_cfg, _pool_failover_active
    with _pool_lock:
        _pool = None
        _pool_cfg = None
        _pool_failover_active = None


def _osc_web_connect():
    """Get a connection from the pool, falling back to direct connect."""
    try:
        pool, cfg = _get_pool()
        conn = pool.get_connection()
        return conn, cfg
    except Exception as exc:
        try:
            mysql_connector, _ = _load_mysql()
            if getattr(getattr(mysql_connector, "errors", object()), "PoolError", None) and isinstance(exc, mysql_connector.errors.PoolError):
                # Pool exhausted – fall back to a direct connection using same cfg
                _, cfg = _get_pool()
                conn = mysql_connector.connect(
                    host=cfg["host"],
                    port=int(cfg["port"]),
                    user=cfg["user"],
                    password=cfg["password"],
                    database=cfg["database"],
                    autocommit=False,
                    charset="utf8mb4",
                    collation="utf8mb4_unicode_ci",
                    connection_timeout=5,
                )
                return conn, cfg
        except Exception:
            pass
        # Pool creation or connection failed – reset and try fresh
        _reset_pool()
        # 讓 failover 模組動態更新 env，確保候選列表反映最新狀態
        try:
            from api.db_failover import probe_remote
            probe_remote()
        except Exception:
            pass
        last_err = None
        for cfg in _osc_web_db_candidates():
            try:
                mysql_connector, _ = _load_mysql()
                conn = mysql_connector.connect(
                    host=cfg["host"],
                    port=int(cfg["port"]),
                    user=cfg["user"],
                    password=cfg["password"],
                    database=cfg["database"],
                    autocommit=False,
                    charset="utf8mb4",
                    collation="utf8mb4_unicode_ci",
                    connection_timeout=5,
                )
                return conn, cfg
            except Exception as e:
                last_err = e
                continue
        raise last_err or RuntimeError("osc_web_db_connect_failed")


# ---------------------------------------------------------------------------
# 2. Core DB execution
# ---------------------------------------------------------------------------

def _osc_exec(sql, params=(), fetch="none"):
    conn, cfg = _osc_web_connect()
    cur = conn.cursor(dictionary=True)
    try:
        cur.execute(sql, params)
        if fetch == "one":
            row = cur.fetchone()
            conn.commit()
            return (_osc_row_json(row) if row else None), cfg
        if fetch == "all":
            rows = cur.fetchall() or []
            conn.commit()
            return [_osc_row_json(r) for r in rows], cfg
        conn.commit()
        return {"rowcount": cur.rowcount, "lastrowid": getattr(cur, "lastrowid", None)}, cfg
    finally:
        try:
            cur.close()
        except Exception:
            logging.getLogger(__name__).debug("silent-catch at %s:%s", __name__, "_osc_exec:cur.close", exc_info=True)
        try:
            conn.close()
        except Exception:
            logging.getLogger(__name__).debug("silent-catch at %s:%s", __name__, "_osc_exec:conn.close", exc_info=True)


# ---------------------------------------------------------------------------
# 3. JSON / path / folder utility functions
# ---------------------------------------------------------------------------

def _osc_json_value(v):
    if isinstance(v, datetime):
        return v.isoformat(sep=" ")
    if isinstance(v, date):
        return v.isoformat()
    if isinstance(v, dt_time):
        return v.isoformat()
    if isinstance(v, timedelta):
        total = int(v.total_seconds())
        sign = "-" if total < 0 else ""
        total = abs(total)
        h = total // 3600
        m = (total % 3600) // 60
        s = total % 60
        return f"{sign}{h:02d}:{m:02d}:{s:02d}"
    if isinstance(v, Decimal):
        return float(v)
    return v


def _osc_row_json(row):
    if not isinstance(row, dict):
        return row
    return {k: _osc_json_value(v) for k, v in row.items()}


def _osc_parse_dt(v):
    if isinstance(v, datetime):
        return v
    if isinstance(v, date):
        return datetime(v.year, v.month, v.day)
    s = str(v or "").strip()
    if not s:
        return None
    s = s.replace("T", " ")
    fmts = [
        "%Y-%m-%d %H:%M:%S",
        "%Y-%m-%d %H:%M",
        "%Y-%m-%d",
        "%Y/%m/%d %H:%M:%S",
        "%Y/%m/%d %H:%M",
        "%Y/%m/%d",
    ]
    for f in fmts:
        try:
            return datetime.strptime(s, f)
        except Exception:
            continue
    return None


def _osc_norm_path(path_str: str) -> str:
    s = str(path_str or "").strip()
    if not s:
        return s
    if s.startswith("/Users/") or s.startswith("/Volumes/") or s.startswith("smb://"):
        return s
    s2 = s.replace("/", "\\")
    up = s2.upper()
    if up.startswith("K:\\SYNOLOGYDRIVE"):
        target = _osc_canonical_active_share_windows()
        return s2.replace("K:\\SynologyDrive", target).replace("K:\\SYNOLOGYDRIVE", target)
    legacy_closed_prefixes = tuple(
        f"K:\\{alias}".upper() for alias in _osc_closed_share_aliases()
    )
    if up.startswith(legacy_closed_prefixes):
        return "Z:" + s2[2:]
    if up.startswith("K:"):
        return "Z:" + s2[2:]
    return s2


def _osc_live_mount_path_candidates(path_str: str) -> list[str]:
    norm = _osc_norm_path(path_str).replace("\\", "/").rstrip("/")
    if not norm:
        return []
    candidates: list[str] = []

    active_prefixes = [
        _osc_canonical_active_share_posix().rstrip("/") + "/01_案件/",
        f"Z:/{_OSC_NAS_HOME_USER}/01_案件/",
        "Z:/home/01_案件/",
    ]
    for prefix in active_prefixes:
        if norm.lower().startswith(prefix.lower()):
            rel = norm[len(prefix):].lstrip("/")
            candidates.extend([
                str(Path.home() / ".magi_mounts" / "homes" / _OSC_NAS_HOME_USER / "01_案件" / rel),
                f"/Volumes/homes/{_OSC_NAS_HOME_USER}/01_案件/{rel}",
            ])

    for alias in _osc_closed_share_aliases():
        for prefix in [
            f"Y:/{alias}/03_工作資料/10_結案/",
            f"/Volumes/{alias}/{alias}/03_工作資料/10_結案/",
        ]:
            if not norm.lower().startswith(prefix.lower()):
                continue
            rel = norm[len(prefix):].lstrip("/")
            candidates.append(str(Path.home() / ".magi_mounts" / alias / alias / "03_工作資料" / "10_結案" / rel))
            for closed_alias in _osc_closed_share_aliases():
                candidates.append(f"/Volumes/{closed_alias}/{closed_alias}/03_工作資料/10_結案/{rel}")
    return [p for p in candidates if p]


def _osc_preferred_smb_local_candidates(path_str: str) -> list[str]:
    """Return live SMB/mount candidates first; tests and runtime may monkeypatch this."""
    return _osc_live_mount_path_candidates(path_str)


def _osc_prioritize_nas_path_candidates(candidates: list[str]) -> list[str]:
    def priority(path: str) -> tuple[int, str]:
        norm = str(path or "").replace("\\", "/")
        if "/.magi_mounts/" in norm:
            return (0, norm)
        if norm.startswith("/Volumes/"):
            return (1, norm)
        if "/Library/CloudStorage/SynologyDrive" in norm or "/SynologyDrive/" in norm:
            return (3, norm)
        if norm.startswith("/Users/"):
            return (2, norm)
        return (4, norm)

    seen: set[str] = set()
    unique: list[str] = []
    for item in candidates:
        text = str(item or "").strip()
        if text and text not in seen:
            seen.add(text)
            unique.append(text)
    return sorted(unique, key=priority)


def _osc_local_path_candidates(path_str: str) -> list[str]:
    """
    Convert legacy Windows/NAS paths into local paths.

    OSC web browsing should prefer live NAS mounts over Synology Drive
    CloudStorage placeholders, because File Provider paths can block for tens
    of seconds before falling back.
    """
    norm = _osc_norm_path(path_str)
    candidates = _osc_preferred_smb_local_candidates(norm) + local_synology_path_candidates(norm)
    return _osc_prioritize_nas_path_candidates(candidates)


def _osc_allowed_local_roots() -> list[str]:
    offline_fixture_root = str(
        os.environ.get("PAPERCLIP_FILEMANAGER_TEST_BASE") or ""
    ).strip()
    if (
        os.environ.get("MAGI_V3_OFFLINE_CERTIFICATION") == "1"
        and offline_fixture_root
    ):
        # Formal certification must be hermetic.  Do not enumerate production
        # NAS/CloudStorage roots (or trigger mount-on-demand) when the
        # certifier has supplied an explicit synthetic file-manager root.
        return [os.path.realpath(offline_fixture_root)]

    def _project_root_markers(root: Path) -> bool:
        try:
            if not root.is_dir():
                return False
        except Exception:
            return False
        return bool((root / "api" / "server.py").exists() and (root / "templates").is_dir())

    def _project_root_candidates() -> list[Path]:
        candidates: list[Path] = []
        seen: set[str] = set()
        try:
            paths = [Path.cwd()]
            paths.extend(list(Path.cwd().parents)[:6])
        except Exception:
            paths = []
        for path in paths:
            try:
                real_path = path.resolve()
            except Exception:
                continue
            key = str(real_path)
            if key in seen:
                continue
            seen.add(key)
            if _project_root_markers(real_path):
                candidates.append(real_path)
        return candidates

    magi_root = Path(__file__).resolve().parents[2]
    magi_root_dir = get_magi_root_dir()
    template_roots = [
        os.environ.get("MAGI_OSC_TEMPLATE_FOLDER", ""),
        str(Path.home() / "Desktop" / "0000-0000-範本-消費者債務清理"),
    ]
    file_manager_test_roots = [
        os.environ.get("PAPERCLIP_FILEMANAGER_TEST_BASE", ""),
        "/tmp/paperclip_filemanager_test_base",
    ]
    closed_alias_roots: list[str] = []
    for share_name in _osc_closed_share_aliases():
        closed_alias_roots.extend([
            str(Path.home() / "Library/CloudStorage/SynologyDrive-homes" / share_name),
            str(Path.home() / "SynologyDrive/homes" / share_name),
            str(Path.home() / "SynologyDrive" / share_name),
            f"/Volumes/{share_name}/{share_name}",
            f"/Volumes/{share_name}-1/{share_name}",
            f"/Volumes/{share_name}-2/{share_name}",
        ])
    magi_mount_roots = []
    magi_mount_home_roots = [
        Path.home() / ".magi_mounts" / "homes" / _OSC_NAS_HOME_USER,
    ]
    for share_name in _osc_closed_share_aliases():
        magi_mount_home_roots.append(Path.home() / ".magi_mounts" / share_name / share_name)
    magi_mount_roots.extend(str(p) for p in magi_mount_home_roots)
    project_roots = [magi_root]
    if magi_root_dir:
        project_roots.append(magi_root_dir)
    project_roots.extend(_project_root_candidates())
    roots = default_synology_share_roots(include_closed=True)
    for env_name in (
        "MAGI_ACTIVE_CASE_ROOT",
        "MAGI_V3_CASE_ROOT",
        "MAGI_CLOSED_CASE_ROOT",
        "MAGI_V3_ARCHIVE_ROOT",
    ):
        configured = str(os.environ.get(env_name) or "").strip()
        if configured:
            roots.append(str(Path(configured).expanduser()))
    for project_root in project_roots:
        if not project_root:
            continue
        roots.extend([
            str(project_root / "exports"),
            str(project_root / "static" / "exports"),
        ])
    roots.extend([
        f"/Volumes/homes/{_OSC_NAS_HOME_USER}",
    ])
    roots += closed_alias_roots + magi_mount_roots + [
        p for p in (template_roots + file_manager_test_roots) if str(p or "").strip()
    ]
    out = []
    for root in roots:
        rp = os.path.realpath(root)
        if rp not in out:
            out.append(rp)
    return out


def _osc_stat_quick(path_str: str, timeout: float = 1.5):
    """os.stat with a timeout guard for stale SMB /Volumes paths."""
    p = str(path_str or "")
    if not _osc_path_needs_fs_timeout(p):
        return os.stat(p)
    result = {"stat": None, "error": None}

    def _run():
        try:
            result["stat"] = os.stat(p)
        except Exception as exc:
            result["error"] = exc

    thread = threading.Thread(target=_run, daemon=True)
    thread.start()
    thread.join(timeout=timeout)
    if thread.is_alive():
        raise TimeoutError(f"stat timeout: {p}")
    if result["error"] is not None:
        raise result["error"]
    return result["stat"]


def _osc_path_needs_fs_timeout(path_str: str) -> bool:
    """True for NAS/File Provider paths that may hang inside direct filesystem calls."""
    norm = str(path_str or "").replace("\\", "/")
    return (
        norm.startswith("/Volumes/")
        or "/.magi_mounts/" in norm
        or "/Library/CloudStorage/SynologyDrive" in norm
        or "/SynologyDrive/" in norm
    )


def _osc_finder_list_folder(path_str: str) -> list[str] | None:
    """Best-effort macOS Finder/Spotlight hook placeholder.

    Kept as a named seam because tests and operators may disable it. Returning
    None means the normal subprocess fallback should be used.
    """
    return None


def _osc_listdir_via_subprocess(path_str: str, timeout: float = 7.0) -> list[str]:
    code = "import json, os, sys\nprint(json.dumps(os.listdir(sys.argv[1]), ensure_ascii=False))"
    try:
        result = subprocess.run(
            [sys.executable, "-c", code, str(path_str)],
            capture_output=True,
            text=True,
            timeout=max(float(timeout), 0.1),
        )
    except subprocess.TimeoutExpired as exc:
        raise TimeoutError(f"listdir timeout: {path_str}") from exc
    if result.returncode != 0:
        stderr = (result.stderr or "").strip()
        raise OSError(stderr.splitlines()[-1] if stderr else "listdir subprocess failed")
    raw = (result.stdout or "").strip()
    if not raw:
        return []
    try:
        data = json.loads(raw)
    except Exception as exc:
        raise OSError(f"listdir subprocess parse failed: {exc}") from exc
    if not isinstance(data, list):
        raise OSError("listdir subprocess returned non-list")
    return [str(x) for x in data]


def _osc_listdir_quick(path_str: str, timeout: float = 1.5) -> list[str]:
    """List a folder with a timeout guard for stale NAS/File Provider paths."""
    p = str(path_str or "")
    if not _osc_path_needs_fs_timeout(p):
        return os.listdir(p)
    finder = _osc_finder_list_folder(p)
    if finder is not None:
        return [str(x) for x in finder]
    result: dict[str, object] = {"names": None, "error": None}

    def _run() -> None:
        try:
            result["names"] = os.listdir(p)
        except Exception as exc:
            result["error"] = exc

    thread = threading.Thread(target=_run, daemon=True)
    thread.start()
    thread.join(max(float(timeout), 0.01))
    if thread.is_alive():
        return _osc_listdir_via_subprocess(p, timeout=max(7.0, float(timeout) * 2))
    if result["error"] is not None:
        raise result["error"]  # type: ignore[misc]
    names = result["names"]
    if isinstance(names, list):
        return [str(x) for x in names]
    return []


def _osc_exists_quick(path_str: str) -> bool:
    try:
        _osc_stat_quick(path_str)
        return True
    except Exception:
        return False


def _osc_isdir_quick(path_str: str) -> bool:
    try:
        return bool(_osc_stat_quick(path_str).st_mode & 0o40000)
    except Exception:
        return False


def _osc_isfile_quick(path_str: str) -> bool:
    try:
        return bool(_osc_stat_quick(path_str).st_mode & 0o100000)
    except Exception:
        return False


def _osc_path_needs_fs_timeout(path_str: str) -> bool:
    """Detect mounts that are usually safer to query through subprocess/timeouts."""
    path = str(path_str or "").replace("\\", "/")
    if not path:
        return False
    lowered = path.lower()
    return (
        "/.magi_mounts/" in lowered
        or lowered.startswith("/volumes/")
        or "/synologydrive" in lowered
        or "/library/cloudstorage/synologydrive" in lowered
    )


def _osc_finder_list_folder(path: str, timeout: float = 0.5) -> list[str] | None:
    """Legacy helper hook for finder/folder listing extension points.

    Returns None when caller should use os.listdir fallback.
    """
    if not _osc_path_needs_fs_timeout(path):
        return None
    return None


def _osc_listdir_with_timeout_subprocess(path: str, timeout: float) -> list[str]:
    code = "import os,sys,json\nprint(json.dumps(os.listdir(sys.argv[1]), ensure_ascii=False))"
    try:
        result = subprocess.run(
            [sys.executable, "-c", code, path],
            capture_output=True,
            text=True,
            timeout=timeout,
        )
    except subprocess.TimeoutExpired as exc:
        raise TimeoutError(f"listdir timed out: {timeout}s") from exc
    if result.returncode != 0:
        message = (result.stderr or "listdir failed").strip().splitlines()[-1] if (result.stderr or "listdir failed") else "listdir failed"
        err_no = 0
        m = re.search(r"Errno\s+(\d+)", message)
        if m:
            try:
                err_no = int(m.group(1))
            except Exception:
                err_no = 0
        if err_no:
            raise OSError(err_no, message)
        raise OSError(message)
    output = (result.stdout or "").strip()
    if not output:
        return []
    try:
        payload = json.loads(output)
    except Exception as exc:
        raise OSError(f"listdir parse failed: {exc}") from exc
    if not isinstance(payload, list):
        raise OSError("listdir helper returned non-list")
    return [str(item) for item in payload]


def _osc_listdir_quick(path: str, timeout: float = 0.25) -> list[str]:
    """Quick directory scan with timeout-aware fallback.

    Legacy-compatible helper used by tests and older callers.
    """
    target = str(path or "")
    if not target:
        return []
    if _osc_finder_list_folder(target) is not None:
        found = _osc_finder_list_folder(target)
        return list(found) if found is not None else []
    if _osc_path_needs_fs_timeout(target):
        try:
            return _osc_listdir_with_timeout_subprocess(target, timeout=timeout)
        except (OSError, TimeoutError):
            pass
    return list(os.listdir(target))


def _osc_is_safe_local_path(path_str: str, *, allow_missing: bool = False) -> bool:
    p = str(path_str or "").strip()
    if not p:
        return False
    try:
        real = os.path.realpath(p)
    except Exception:
        return False
    for root in _osc_allowed_local_roots():
        if real == root or real.startswith(root + os.sep):
            return bool(allow_missing or _osc_exists_quick(real))
    return False


def _osc_request_path_is_allowed(path_str: str) -> bool:
    """Authorize an OSC request path before testing whether it exists.

    Missing paths below approved roots remain distinguishable from traversal
    attempts. Realpath resolution in ``_osc_is_safe_local_path`` also prevents
    symlinks below an approved root from escaping it.
    """

    raw = str(path_str or "").strip()
    if not raw:
        return False
    candidates = _osc_local_path_candidates(raw)
    norm = _osc_norm_path(raw).replace("\\", "/")
    if norm and norm not in candidates:
        candidates.append(norm)
    return any(
        _osc_is_safe_local_path(candidate, allow_missing=True)
        for candidate in candidates
    )


def _osc_resolve_existing_local_path(path_str: str, *, prefer_dir: Optional[bool] = None) -> str:
    candidates = _osc_local_path_candidates(path_str)
    norm = _osc_norm_path(path_str).replace("\\", "/")
    if norm and norm not in candidates:
        candidates.append(norm)
    for cand in candidates:
        try:
            real = os.path.realpath(cand)
            if not _osc_is_safe_local_path(real):
                continue
            if not _osc_exists_quick(real):
                continue
            if prefer_dir is True and not _osc_isdir_quick(real):
                continue
            if prefer_dir is False and not _osc_isfile_quick(real):
                continue
            return real
        except Exception:
            continue
    return ""


def _osc_relpath_under(base_path: str, target_path: str) -> str:
    try:
        rel = os.path.relpath(os.path.realpath(target_path), os.path.realpath(base_path))
    except Exception:
        return ""
    if rel in {".", ""}:
        return ""
    return rel.replace("\\", "/")


_OSC_PATH_REFERENCE_COLUMNS = (
    ("cases", "folder_path"),
    ("document_index", "file_path"),
    ("case_documents", "file_path"),
    ("case_todos", "source_file"),
    ("legal_insights", "source_file"),
    ("court_judgments", "source_file"),
)


def _osc_replace_path_prefix_references(old_path: str, new_path: str, *, exec_fn=None) -> dict:
    """Replace known OSC path prefixes after a filesystem rename.

    Paths may be stored either as canonical Windows/NAS paths or as local macOS
    Synology paths. The caller should pass the concrete old/new pair it knows;
    this helper also tries the canonical mapping so document indexes and todo
    source paths do not keep pointing at the renamed folder.
    """
    old_raw = str(old_path or "").strip()
    new_raw = str(new_path or "").strip()
    if not old_raw or not new_raw or old_raw == new_raw:
        return {"updated": 0, "attempted": 0, "errors": []}

    pairs: list[tuple[str, str]] = []

    def _add_pair(old_value: str, new_value: str) -> None:
        old_value = str(old_value or "").strip().rstrip("/\\")
        new_value = str(new_value or "").strip().rstrip("/\\")
        if old_value and new_value and old_value != new_value and (old_value, new_value) not in pairs:
            pairs.append((old_value, new_value))

    _add_pair(old_raw, new_raw)
    try:
        _add_pair(translate_local_path_to_canonical(old_raw), translate_local_path_to_canonical(new_raw))
    except Exception:
        logger.debug("silent-catch canonical path pair for rename", exc_info=True)
    try:
        _add_pair(_osc_norm_path(old_raw), _osc_norm_path(new_raw))
    except Exception:
        logger.debug("silent-catch normalized path pair for rename", exc_info=True)

    exec_fn = exec_fn or _osc_exec
    updated = 0
    attempted = 0
    errors: list[str] = []
    for old_value, new_value in pairs:
        old_len = len(old_value)
        for table, column in _OSC_PATH_REFERENCE_COLUMNS:
            attempted += 1
            sql = f"""
                UPDATE `{table}`
                   SET `{column}` = CONCAT(%s, SUBSTRING(`{column}`, %s))
                 WHERE `{column}` = %s
                    OR (
                        LEFT(`{column}`, %s) = %s
                        AND (
                            SUBSTRING(`{column}`, %s, 1) = '/'
                            OR SUBSTRING(`{column}`, %s, 1) = CHAR(92)
                        )
                    )
            """
            params = (new_value, old_len + 1, old_value, old_len, old_value, old_len + 1, old_len + 1)
            try:
                result, _ = exec_fn(sql, params, fetch="none")
                if isinstance(result, dict):
                    updated += int(result.get("rowcount") or 0)
            except Exception as exc:
                # Some installs may not have every optional table/column yet.
                msg = str(exc)
                if (
                    "Unknown column" in msg
                    or "doesn't exist" in msg
                    or "does not exist" in msg
                    or "no such table" in msg.lower()
                    or "no such column" in msg.lower()
                ):
                    logger.debug("silent-catch missing optional path reference %s.%s", table, column, exc_info=True)
                    continue
                errors.append(f"{table}.{column}: {exc}")
                logger.debug("silent-catch replace path reference %s.%s", table, column, exc_info=True)
    return {"updated": updated, "attempted": attempted, "errors": errors[:8]}


def _osc_human_size(size: int) -> str:
    units = ["B", "KB", "MB", "GB", "TB"]
    n = float(max(0, int(size or 0)))
    unit = units[0]
    for unit in units:
        if n < 1024 or unit == units[-1]:
            break
        n /= 1024.0
    return f"{n:.1f}{unit}" if unit != "B" else f"{int(n)}B"


_OSC_HIDDEN_FILE_PATTERNS = (
    re.compile(r"^\.DS_Store$"),
    re.compile(r"^Thumbs\.db$", re.IGNORECASE),
    re.compile(r"^~\$.*"),
    re.compile(r"^\._.*"),
    re.compile(r"^\.synology.*", re.IGNORECASE),
    re.compile(r"^\.DocumentRevisions.*", re.IGNORECASE),
    re.compile(r"^.*\.tmp$", re.IGNORECASE),
    re.compile(r"^\.Spotlight.*"),
    re.compile(r"^\.Trashes$"),
    re.compile(r"^\.fseventsd$"),
)


def _osc_is_hidden_folder_entry(name: str) -> bool:
    return any(pattern.match(name or "") for pattern in _OSC_HIDDEN_FILE_PATTERNS)


def _osc_directory_metadata_quick(path_str: str, timeout: float = 1.5) -> list[dict]:
    """List one directory level with bounded I/O on NAS/File Provider paths."""
    path = str(path_str or "")
    if not _osc_path_needs_fs_timeout(path):
        rows: list[dict] = []
        with os.scandir(path) as scan:
            for entry in scan:
                try:
                    item_stat = entry.stat(follow_symlinks=False)
                except OSError:
                    continue
                is_dir = entry.is_dir(follow_symlinks=False)
                rows.append(
                    {
                        "name": entry.name,
                        "is_dir": is_dir,
                        "size": 0 if is_dir else int(item_stat.st_size),
                        "mtime": float(item_stat.st_mtime),
                    }
                )
        return rows

    code = (
        "import json,os,stat,sys\n"
        "out=[]\n"
        "with os.scandir(sys.argv[1]) as scan:\n"
        " for entry in scan:\n"
        "  try:\n"
        "   st=entry.stat(follow_symlinks=False)\n"
        "   is_dir=stat.S_ISDIR(st.st_mode)\n"
        "   out.append({'name':entry.name,'is_dir':is_dir,'size':0 if is_dir else int(st.st_size),'mtime':float(st.st_mtime)})\n"
        "  except OSError:\n"
        "   pass\n"
        "print(json.dumps(out,ensure_ascii=False))\n"
    )
    try:
        result = subprocess.run(
            [sys.executable, "-c", code, path],
            capture_output=True,
            text=True,
            timeout=max(0.1, float(timeout)),
        )
    except subprocess.TimeoutExpired as exc:
        raise TimeoutError(f"folder metadata timed out: {float(timeout):.1f}s") from exc
    if result.returncode != 0:
        message = (result.stderr or "folder metadata failed").strip()
        raise OSError(message.splitlines()[-1] if message else "folder metadata failed")
    try:
        payload = json.loads((result.stdout or "[]").strip() or "[]")
    except Exception as exc:
        raise OSError(f"folder metadata parse failed: {exc}") from exc
    if not isinstance(payload, list):
        raise OSError("folder metadata returned non-list")
    return [row for row in payload if isinstance(row, dict) and row.get("name") is not None]


def _osc_folder_entries(base_path: str, relative_path: str = "", limit: int = 240) -> dict:
    base_real = os.path.realpath(base_path)
    target_real = os.path.realpath(os.path.join(base_real, relative_path or ""))
    if not _osc_is_safe_local_path(base_real):
        return {"ok": False, "error": "base_not_allowed"}
    if target_real != base_real and not target_real.startswith(base_real + os.sep):
        return {"ok": False, "error": "path_escape"}
    if not _osc_isdir_quick(target_real):
        return {"ok": False, "error": "folder_not_found"}
    entries = []
    try:
        metadata = _osc_directory_metadata_quick(target_real)
    except Exception as e:
        return {"ok": False, "error": str(e)}
    visible_rows = [
        row for row in metadata
        if not _osc_is_hidden_folder_entry(str(row.get("name") or ""))
    ]
    visible_rows.sort(
        key=lambda row: (
            not bool(row.get("is_dir")),
            str(row.get("name") or "").lower(),
        )
    )
    for row in visible_rows[:max(1, int(limit or 240))]:
        name = str(row.get("name") or "")
        if not name:
            continue
        full = os.path.join(target_real, name)
        is_dir = bool(row.get("is_dir"))
        size = int(row.get("size") or 0)
        modified = float(row.get("mtime") or 0)
        rel = _osc_relpath_under(base_real, full)
        entries.append(
            {
                "name": name,
                "relative_path": rel,
                "type": "dir" if is_dir else "file",
                "ext": "" if is_dir else os.path.splitext(name)[1].lower(),
                "size": None if is_dir else size,
                "size_label": "" if is_dir else _osc_human_size(size),
                "modified_at": datetime.fromtimestamp(modified).strftime("%Y-%m-%d %H:%M:%S") if modified else "",
            }
        )
    parent_relative = _osc_relpath_under(base_real, os.path.dirname(target_real)) if target_real != base_real else ""
    return {
        "ok": True,
        "base_path": base_real,
        "current_path": target_real,
        "current_relative_path": _osc_relpath_under(base_real, target_real),
        "parent_relative_path": parent_relative,
        "entries": entries,
    }


# ---------------------------------------------------------------------------
# 4. File reading utilities
# ---------------------------------------------------------------------------

_OSC_TEXT_EXTENSIONS = {
    ".txt", ".md", ".json", ".csv", ".tsv", ".yaml", ".yml",
    ".xml", ".html", ".htm", ".log", ".py", ".js", ".ts", ".css",
}


def _osc_is_editable_text_path(path_str: str) -> bool:
    ext = os.path.splitext(str(path_str or "").strip())[1].lower()
    return ext in _OSC_TEXT_EXTENSIONS


def _osc_read_text_file(path_str: str, max_bytes: int = 2 * 1024 * 1024) -> tuple[str, str]:
    size = os.path.getsize(path_str)
    if size > max_bytes:
        raise ValueError(f"file_too_large:{size}")
    # SMB EDEADLK retry：macOS SMB-over-Tailscale 偶發 errno 11 / 35
    import time as _time
    raw = b""
    last_exc: Exception | None = None
    for attempt in range(4):
        try:
            raw = Path(path_str).read_bytes()
            last_exc = None
            break
        except OSError as e:
            last_exc = e
            if e.errno in (11, 35) and attempt < 3:
                _time.sleep(0.25 * (2 ** attempt))
                continue
            raise
    if last_exc:
        raise last_exc
    for enc in ("utf-8-sig", "utf-8", "cp950"):
        try:
            return raw.decode(enc), enc
        except Exception:
            continue
    return raw.decode("utf-8", errors="replace"), "utf-8-replace"


def _osc_smb_candidates(path_str: str) -> list[str]:
    """
    Return ordered SMB candidate URLs for NAS.
    """
    try:
        from api.nas_mount_guard import resolve_nas_host
        host = resolve_nas_host()
    except Exception:
        try:
            from api.routing.node_registry import get_node as _get_nas_node
            _nas = _get_nas_node("nas")
            host = (os.environ.get("MAGI_NAS_HOST") or (_nas.lan_ip if _nas else None) or "").strip()
        except Exception:
            host = (os.environ.get("MAGI_NAS_HOST") or "").strip()
    p = _osc_norm_path(path_str).replace("\\", "/")
    if p.startswith("/Users/") or p.startswith("/Volumes/"):
        p = translate_local_path_to_canonical(p).replace("\\", "/")
    out: list[str] = []
    rel = ""
    active_prefix = _osc_canonical_active_share_posix()
    if p.startswith(active_prefix):
        rel = p[len(active_prefix):].lstrip("/")
        for base in [f"smb://{host}/SynologyDrive", f"smb://{host}/home", f"smb://{host}/homes/{_OSC_NAS_HOME_USER}"]:
            out.append(f"{base}/{rel}" if rel else base)
    elif p.startswith("Y:/"):
        rel = p[len("Y:/"):].lstrip("/")
        closed_prefix = _OSC_NAS_CLOSED_SHARE_NAME + "/"
        if rel.lower().startswith(closed_prefix.lower()):
            rel = rel[len(closed_prefix):]
        for base in [f"smb://{host}/{_OSC_NAS_CLOSED_SHARE_NAME}/{_OSC_NAS_CLOSED_SHARE_NAME}", f"smb://{host}/{_OSC_NAS_CLOSED_SHARE_NAME}", f"smb://{host}/home"]:
            out.append(f"{base}/{rel}" if rel else base)
    elif p.lower().startswith("smb://"):
        out.append(p)
    elif re.match(r"^[A-Za-z]:/", p):
        rel = p[3:].lstrip("/")
        out.append(f"smb://{host}/{rel}" if rel else f"smb://{host}")
    else:
        out.append(f"smb://{host}")

    # always keep NAS root as last resort
    out.append(f"smb://{host}")
    seen = set()
    uniq = []
    for x in out:
        if x not in seen:
            seen.add(x)
            uniq.append(x)
    return uniq


def _osc_path_to_smb(path_str: str) -> str:
    """
    Convert canonical Windows path to smb:// URL for NAS browsing.
    """
    cands = _osc_smb_candidates(path_str)
    return cands[0] if cands else str(path_str or "")


def _osc_windows_unc_candidates(path_str: str) -> list[str]:
    """
    Return Windows UNC \\\\nas-host\\share\\... candidates for NAS browsing on Win.

    Maps configured Y:/archive and Z:/active-share paths back to NAS shares.
    Used by web UI on Windows clients (Explorer / file:// fallback).
    """
    try:
        from api.nas_mount_guard import resolve_nas_host
        host = resolve_nas_host()
    except Exception:
        host = (os.environ.get("MAGI_NAS_HOST") or "").strip() or ""
    candidates: list[str] = []
    norm = _osc_norm_path(path_str)
    np = norm.replace("/", "\\")
    # Y:\<archive-share>\... → \\nas-host\<archive-share>\...
    if np.upper().startswith("Y:\\"):
        rel = np[3:]  # 去掉 Y:\
        closed_prefix = _OSC_NAS_CLOSED_SHARE_NAME + "\\"
        if rel.lower().startswith(closed_prefix.lower()):
            sub = rel[len(closed_prefix):]
            candidates.append(f"\\\\{host}\\{_OSC_NAS_CLOSED_SHARE_NAME}\\{sub}")
            candidates.append(f"\\\\{host}\\{_OSC_NAS_CLOSED_SHARE_NAME}\\{_OSC_NAS_CLOSED_SHARE_NAME}\\{sub}")
        else:
            candidates.append(f"\\\\{host}\\{rel}")
    # Z:\<active-share>\... → \\nas-host\homes\<user>\... + \\nas-host\SynologyDrive\...
    elif np.upper().startswith("Z:\\"):
        rel = np[3:]
        home_prefix = _OSC_NAS_HOME_USER + "\\"
        if rel.lower().startswith(home_prefix.lower()):
            sub = rel[len(home_prefix):]
            candidates.append(f"\\\\{host}\\homes\\{_OSC_NAS_HOME_USER}\\{sub}")
            candidates.append(f"\\\\{host}\\SynologyDrive\\{sub}")
            candidates.append(f"\\\\{host}\\home\\{sub}")
        else:
            candidates.append(f"\\\\{host}\\{rel}")
    # 已是 \\... UNC 直接保留
    elif np.startswith("\\\\"):
        candidates.append(np)
    # dedup keep order
    seen = set()
    uniq = []
    for c in candidates:
        if c and c not in seen:
            seen.add(c)
            uniq.append(c)
    return uniq


def _osc_windows_synology_candidates(path_str: str) -> list[str]:
    """
    Return Windows Synology Drive default-install paths (C:\\Users\\<user>\\SynologyDrive\\...).

    Tries to extract the case-relative subpath using common markers
    (01_案件 / 02_開辦 / 03_閱卷 / 04_我方) and reconstruct under SynologyDrive.
    """
    candidates: list[str] = []
    norm = _osc_norm_path(path_str)
    np = norm.replace("/", "\\")
    markers = ["01_案件", "02_開辦", "03_閱卷", "04_我方"]
    for marker in markers:
        if marker in np:
            # 取 marker 開始的相對路徑
            idx = np.find(marker)
            rel = np[idx:]  # e.g. "01_案件\foo\bar"
            candidates.append(f"%USERPROFILE%\\SynologyDrive\\{rel}")
            candidates.append(f"C:\\SynologyDrive\\{rel}")
            break
    # dedup
    seen = set()
    uniq = []
    for c in candidates:
        if c and c not in seen:
            seen.add(c)
            uniq.append(c)
    return uniq


def _osc_try_open_path(path_str: str) -> dict:
    """
    Best-effort open folder in host OS. Returns execution result only.
    """
    p = str(path_str or "").strip()
    if not p:
        return {"ok": False, "error": "empty_path"}
    try:
        if platform.system() == "Darwin":
            subprocess.run(["open", p], stdout=subprocess.DEVNULL, stderr=subprocess.DEVNULL, timeout=10)
            return {"ok": True, "method": "open"}
        if platform.system().lower().startswith("win"):
            os.startfile(p)  # type: ignore[attr-defined]
            return {"ok": True, "method": "startfile"}
        subprocess.run(["xdg-open", p], stdout=subprocess.DEVNULL, stderr=subprocess.DEVNULL, timeout=10)
        return {"ok": True, "method": "xdg-open"}
    except Exception as e:
        return {"ok": False, "error": str(e)}


def _osc_case_folder_from_doc_path(case_number: str, file_path: str) -> str:
    """
    Best-effort derive case folder from a document path.
    """
    p = str(file_path or "").strip().replace("\\", "/")
    if not p:
        return ""
    cn = str(case_number or "").strip()
    parts = [x for x in p.split("/") if x]
    if cn and cn in parts:
        idx = parts.index(cn)
        return "/" + "/".join(parts[: idx + 1])
    if cn:
        for i, seg in enumerate(parts):
            if cn in seg:
                return "/" + "/".join(parts[: i + 1])
    return str(Path(p).parent)


def _osc_guess_case_folder(case_number: str) -> str:
    """
    Fallback when cases.folder_path is empty:
    infer from document_index/case_documents latest file path.
    """
    cn = (case_number or "").strip()
    if not cn:
        return ""
    try:
        row, _ = _osc_exec(
            """
            SELECT file_path
            FROM document_index
            WHERE case_number=%s AND file_path IS NOT NULL AND file_path != ''
            ORDER BY modified_date DESC, id DESC
            LIMIT 1
            """,
            (cn,),
            fetch="one",
        )
        if row and row.get("file_path"):
            return _osc_case_folder_from_doc_path(cn, row.get("file_path"))
    except Exception:
        logging.getLogger(__name__).debug("silent-catch at %s:%s", __name__, "_osc_guess_case_folder:doc_index", exc_info=True)
    try:
        row, _ = _osc_exec(
            """
            SELECT file_path
            FROM case_documents
            WHERE file_path IS NOT NULL AND file_path != '' AND (description LIKE %s OR file_name LIKE %s)
            ORDER BY upload_date DESC, id DESC
            LIMIT 1
            """,
            (f"%{cn}%", f"%{cn}%"),
            fetch="one",
        )
        if row and row.get("file_path"):
            return _osc_case_folder_from_doc_path(cn, row.get("file_path"))
    except Exception:
        logging.getLogger(__name__).debug("silent-catch at %s:%s", __name__, "_osc_guess_case_folder:case_docs", exc_info=True)
    # 最後手段：掃 NAS 01_案件/ 找含 case_number 或 client_name 的資料夾
    # 受 CLAUDE.md §4.6 NAS I/O 節流規範：限深度 5、上限 2000、每 50 個 sleep
    try:
        scanned = _osc_scan_nas_for_case_folder(cn)
        if scanned:
            return scanned
    except Exception:
        logging.getLogger(__name__).debug("silent-catch at %s:%s", __name__, "_osc_guess_case_folder:nas_scan", exc_info=True)
    return ""


def _osc_scan_nas_for_case_folder(case_number: str, *, client_name: str = "") -> str:
    """掃 NAS 01_案件/ 找符合 case_number 或 client_name 的案件資料夾。

    為避免一次掃整個 NAS（CLAUDE.md §4.6），限制 depth=5 / max 2000 / 每 50 個
    sleep 0.05s。回傳 canonical 路徑（Y:\\lumi\\01_案件\\... 形式），找不到回 ""。
    """
    cn = (case_number or "").strip()
    cln = (client_name or "").strip()
    if not cn and not cln:
        return ""
    import time as _time
    tokens = [t for t in [cn, cln] if t]
    # 找 NAS 01_案件 base：枚舉允許的 root。allowed roots 有兩種常見型態：
    #   /Volumes/homes/<account>                    -> 直接含 01_案件
    #   /Users/.../SynologyDrive-homes/<share/user> -> 需往下一層找 01_案件
    for root in _osc_allowed_local_roots():
        try:
            bases: list[str] = []
            direct = os.path.join(root, "01_案件")
            if os.path.isdir(direct):
                bases.append(direct)
            for top in os.listdir(root):
                if "lumi" not in top.lower():
                    continue
                base = os.path.join(root, top, "01_案件")
                if not os.path.isdir(base):
                    continue
                bases.append(base)
            for base in _osc_unique_strings(bases):
                count = 0
                for dirpath, dirnames, _ in os.walk(base):
                    rel_depth = dirpath[len(base):].count(os.sep)
                    if rel_depth >= 5:
                        dirnames[:] = []
                        continue
                    for d in dirnames:
                        if any(t in d for t in tokens):
                            full = os.path.join(dirpath, d)
                            try:
                                from api.case_path_mapper import translate_local_path_to_canonical
                                return translate_local_path_to_canonical(full)
                            except Exception:
                                return full
                        count += 1
                        if count % 50 == 0:
                            _time.sleep(0.05)
                        if count > 2000:
                            return ""
        except OSError:
            continue
    return ""


# ---------------------------------------------------------------------------
# 5. Fulltext / HTML utilities
# ---------------------------------------------------------------------------

def _osc_strip_html_to_text(html: str) -> str:
    s = html or ""
    s = re.sub(r"(?is)<script[^>]*>.*?</script>", " ", s)
    s = re.sub(r"(?is)<style[^>]*>.*?</style>", " ", s)
    s = re.sub(r"(?is)<noscript[^>]*>.*?</noscript>", " ", s)
    s = re.sub(r"(?is)<br\s*/?>", "\n", s)
    s = re.sub(r"(?is)</p>", "\n", s)
    s = re.sub(r"(?is)<[^>]+>", " ", s)
    s = ihtml.unescape(s)
    s = re.sub(r"[ \t\r\f\v]+", " ", s)
    s = re.sub(r"\n\s*\n+", "\n\n", s)
    return s.strip()


def _osc_fetch_url_text(url: str, timeout: int = 20) -> dict:
    import ssl as _ssl
    u = str(url or "").strip()
    if not u.lower().startswith(("http://", "https://")):
        return {"ok": False, "error": "invalid_url"}
    headers = {
        "User-Agent": "Mozilla/5.0 (Macintosh; Intel Mac OS X 14_0) AppleWebKit/537.36 (KHTML, like Gecko) Chrome/122 Safari/537.36",
        "Accept-Language": "zh-TW,zh;q=0.9,en;q=0.8",
    }
    req = urllib.request.Request(u, headers=headers)
    to = max(5, int(timeout))
    # Try normal SSL first, then fallback to unverified context
    for ctx in (None, _ssl.create_default_context()):
        try:
            if ctx is not None:
                ctx.check_hostname = False
                ctx.verify_mode = _ssl.CERT_NONE
            with urllib.request.urlopen(req, timeout=to, context=ctx) as resp:
                raw = resp.read()
                ct = (resp.headers.get("Content-Type") or "").lower()
            if "charset=" in ct:
                enc = ct.split("charset=", 1)[1].split(";", 1)[0].strip()
            else:
                enc = "utf-8"
            try:
                html_str = raw.decode(enc, errors="ignore")
            except Exception:
                html_str = raw.decode("utf-8", errors="ignore")
            text = _osc_strip_html_to_text(html_str)
            if len(text) < 80:
                return {"ok": False, "error": "content_too_short", "text": text}
            return {"ok": True, "text": text, "html_len": len(html_str)}
        except _ssl.SSLError:
            if ctx is not None:
                return {"ok": False, "error": "ssl_failed_even_unverified"}
            continue  # retry with unverified context
        except urllib.error.HTTPError as e:
            return {"ok": False, "error": f"http_{e.code}"}
        except Exception as e:
            if "ssl" in str(e).lower() or "certificate" in str(e).lower():
                if ctx is not None:
                    return {"ok": False, "error": str(e)}
                continue  # retry with unverified context
            return {"ok": False, "error": str(e)}
    return {"ok": False, "error": "fetch_exhausted"}


def _osc_lookup_fulltext_fallback(title: str = "", case_number: str = "", url: str = "") -> dict:
    """
    Fallback source for sites that block direct scraping (e.g. login-gated pages).
    Try retrieving full text from local DB mirrors first.
    """
    t = (title or "").strip()
    cn = (case_number or "").strip()
    u = (url or "").strip()
    try:
        from api.osc.insight_filters import is_extractive_fast_judgment_digest
    except Exception:
        is_extractive_fast_judgment_digest = lambda *values: False
    try:
        params = []
        clauses = []
        if cn:
            clauses.append("case_number LIKE %s")
            params.append(f"%{cn}%")
        if t:
            clauses.append("(document_name LIKE %s OR case_reason LIKE %s)")
            params.extend([f"%{t[:80]}%", f"%{t[:80]}%"])
        if u:
            clauses.append("source_file LIKE %s")
            params.append(f"%{u}%")
        if clauses:
            row, _ = _osc_exec(
                f"""
                SELECT id, document_name, case_number, case_reason, source_file, insight_text, raw_text
                FROM legal_insights
                WHERE {' OR '.join(clauses)}
                ORDER BY CHAR_LENGTH(COALESCE(insight_text, raw_text, '')) DESC, extracted_date DESC, id DESC
                LIMIT 1
                """,
                tuple(params),
                fetch="one",
            )
            if row:
                txt = (row.get("insight_text") or row.get("raw_text") or "").strip()
                if len(txt) >= 300 and not is_non_extractable_legal_insight(
                    row.get("document_name"),
                    row.get("case_reason"),
                    row.get("source_file"),
                    txt,
                ):
                    return {
                        "ok": True,
                        "text": txt,
                        "source": "fallback_legal_insights",
                        "matched": {
                            "id": row.get("id"),
                            "title": row.get("document_name") or "",
                            "case_number": row.get("case_number") or "",
                        },
                    }
    except Exception:
        logging.getLogger(__name__).debug("silent-catch at %s:%s", __name__, "_osc_lookup_fulltext_fallback:legal_insights", exc_info=True)
    try:
        params = []
        clauses = []
        if cn:
            clauses.append("case_number LIKE %s")
            params.append(f"%{cn}%")
        if t:
            clauses.append("(case_number LIKE %s OR summary LIKE %s)")
            params.extend([f"%{t[:40]}%", f"%{t[:120]}%"])
        if u:
            clauses.append("source_url LIKE %s")
            params.append(f"%{u}%")
        if clauses:
            row, _ = _osc_exec(
                f"""
                SELECT id, court_name, case_number, summary, full_text, source_url
                FROM court_judgments
                WHERE {' OR '.join(clauses)}
                ORDER BY CHAR_LENGTH(COALESCE(full_text, summary, '')) DESC, crawled_at DESC, id DESC
                LIMIT 1
                """,
                tuple(params),
                fetch="one",
            )
            if row:
                summary_text = (row.get("summary") or "").strip()
                if is_extractive_fast_judgment_digest(summary_text) and not (row.get("full_text") or "").strip():
                    return {"ok": False, "error": "fast_digest_without_full_text"}
                txt = (row.get("full_text") or summary_text).strip()
                if len(txt) >= 300:
                    return {
                        "ok": True,
                        "text": txt,
                        "source": "fallback_court_judgments",
                        "matched": {
                            "id": row.get("id"),
                            "title": f"{row.get('court_name') or ''} {row.get('case_number') or ''}".strip(),
                            "case_number": row.get("case_number") or "",
                        },
                    }
    except Exception:
        logging.getLogger(__name__).debug("silent-catch at %s:%s", __name__, "_osc_lookup_fulltext_fallback:court_judgments", exc_info=True)
    return {"ok": False, "error": "fallback_not_found"}


# ---------------------------------------------------------------------------
# 6. Skill execution utilities
# ---------------------------------------------------------------------------

def _osc_run_skill(skill: str, task: str, timeout_sec: int = 180, route_key: str = "") -> dict:
    try:
        from api.routing.service_registry import get_service_url as _get_svc_url
        _tools_default = _get_svc_url("tools_api")
    except Exception:
        _tools_default = "http://127.0.0.1:5003"
    tools_api = (os.environ.get("MAGI_TOOLS_API") or _tools_default).rstrip("/")
    payload = {
        "skill": skill,
        "task": task,
        "timeout_sec": int(timeout_sec),
        "auto_repair": False,
        "rollback_on_fail": True,
        "auto_install_deps": False,
        "route_key": route_key or f"osc:{skill}",
    }
    data = json.dumps(payload, ensure_ascii=False).encode("utf-8")
    req = urllib.request.Request(
        tools_api + "/skills/run",
        data=data,
        headers={"Content-Type": "application/json"},
    )
    try:
        with urllib.request.urlopen(req, timeout=min(90, max(20, int(timeout_sec) + 30))) as resp:
            raw = resp.read().decode("utf-8", errors="replace")
            return json.loads(raw or "{}")
    except urllib.error.HTTPError as e:
        try:
            body = e.read().decode("utf-8", errors="replace")
        except Exception:
            body = ""
        return {"success": False, "error": f"http_{getattr(e, 'code', 0)}", "body": body[:800]}
    except Exception as e:
        return {"success": False, "error": str(e)[:240]}


def _osc_skill_json_task(command: str, payload: dict) -> str:
    return f"{str(command or '').strip()}{json.dumps(payload or {}, ensure_ascii=False, separators=(',', ':'))}"


def _osc_parse_skill_output(run_result: dict) -> dict:
    if (not isinstance(run_result, dict)) or (not run_result.get("success")):
        return {"success": False, "error": (run_result.get("error") if isinstance(run_result, dict) else "run_failed")}
    raw = str(run_result.get("output") or "").strip()
    if not raw:
        return {"success": False, "error": "empty_output"}
    try:
        obj = json.loads(raw)
    except Exception as e:
        return {"success": False, "error": f"output_json_parse_failed: {e}"}
    return obj if isinstance(obj, dict) else {"success": False, "error": "output_not_dict"}


# ---------------------------------------------------------------------------
# 7. Parsing and normalization utilities
# ---------------------------------------------------------------------------

def _osc_title_norm(s: str) -> str:
    x = str(s or "")
    x = x.replace("（", "(").replace("）", ")")
    x = re.sub(r"\s+", "", x)
    x = re.sub(r"[，。、「」『』【】《》:：;；,./\\|_+\-\(\)\[\]{}·]", "", x)
    return x.lower()


_OSC_JUDICIAL_COURT_SEARCH_LABELS = {
    "最高法院": "最高法院",
    "最高行政法院": "最高行政法院(含改制前行政法院)",
    "臺北高等行政法院": "臺北高等行政法院",
    "臺中高等行政法院": "臺中高等行政法院",
    "高雄高等行政法院": "高雄高等行政法院",
    "智慧財產及商業法院": "智慧財產及商業法院",
    "懲戒法院": "懲戒法院",
    "臺灣高等法院": "臺灣高等法院",
    "臺灣高等法院臺中分院": "臺灣高等法院臺中分院",
    "臺灣高等法院臺南分院": "臺灣高等法院臺南分院",
    "臺灣高等法院高雄分院": "臺灣高等法院高雄分院",
    "臺灣高等法院花蓮分院": "臺灣高等法院花蓮分院",
    "臺灣臺北地方法院": "臺灣臺北地方法院",
    "臺灣士林地方法院": "臺灣士林地方法院",
    "臺灣新北地方法院": "臺灣新北地方法院",
    "臺灣桃園地方法院": "臺灣桃園地方法院",
    "臺灣新竹地方法院": "臺灣新竹地方法院",
    "臺灣苗栗地方法院": "臺灣苗栗地方法院",
    "臺灣臺中地方法院": "臺灣臺中地方法院",
    "臺灣南投地方法院": "臺灣南投地方法院",
    "臺灣彰化地方法院": "臺灣彰化地方法院",
    "臺灣雲林地方法院": "臺灣雲林地方法院",
    "臺灣嘉義地方法院": "臺灣嘉義地方法院",
    "臺灣臺南地方法院": "臺灣臺南地方法院",
    "臺灣高雄地方法院": "臺灣高雄地方法院",
    "臺灣橋頭地方法院": "臺灣橋頭地方法院",
    "臺灣屏東地方法院": "臺灣屏東地方法院",
    "臺灣臺東地方法院": "臺灣臺東地方法院",
    "臺灣花蓮地方法院": "臺灣花蓮地方法院",
    "臺灣宜蘭地方法院": "臺灣宜蘭地方法院",
    "臺灣基隆地方法院": "臺灣基隆地方法院",
    "臺灣澎湖地方法院": "臺灣澎湖地方法院",
    "福建金門地方法院": "福建金門地方法院",
    "福建連江地方法院": "福建連江地方法院",
}

_OSC_JUDICIAL_COURT_ALIASES = {
    "台灣高等法院": "臺灣高等法院",
    "高等法院": "臺灣高等法院",
    "台灣高等法院台中分院": "臺灣高等法院臺中分院",
    "台灣高等法院台南分院": "臺灣高等法院臺南分院",
    "台灣高等法院高雄分院": "臺灣高等法院高雄分院",
    "台灣高等法院花蓮分院": "臺灣高等法院花蓮分院",
    "臺灣高等法院台中分院": "臺灣高等法院臺中分院",
    "臺灣高等法院台南分院": "臺灣高等法院臺南分院",
    "高等法院台中分院": "臺灣高等法院臺中分院",
    "高等法院台南分院": "臺灣高等法院臺南分院",
    "高等法院高雄分院": "臺灣高等法院高雄分院",
    "高等法院花蓮分院": "臺灣高等法院花蓮分院",
    "台北高等行政法院": "臺北高等行政法院",
    "台中高等行政法院": "臺中高等行政法院",
    "台灣台北地方法院": "臺灣臺北地方法院",
    "台灣台中地方法院": "臺灣臺中地方法院",
    "台灣台南地方法院": "臺灣臺南地方法院",
    "台灣台東地方法院": "臺灣臺東地方法院",
}


def _osc_unique_keep_order(values: list[str]) -> list[str]:
    seen = set()
    out = []
    for v in values:
        s = str(v or "").strip()
        if (not s) or (s in seen):
            continue
        seen.add(s)
        out.append(s)
    return out


def _osc_normalize_court_name(s: str) -> str:
    text = str(s or "").strip()
    if not text:
        return ""
    text = text.replace("台", "臺")
    return _OSC_JUDICIAL_COURT_ALIASES.get(text, text)


def _osc_extract_court_names(s: str) -> list[str]:
    text = str(s or "").strip()
    if not text:
        return []
    text2 = text.replace("台", "臺")
    hits: list[str] = []

    # Match more specific court names first to avoid a generic "臺灣高等法院"
    # hit swallowing a branch court name.
    for court_name in sorted(_OSC_JUDICIAL_COURT_SEARCH_LABELS.keys(), key=len, reverse=True):
        if court_name and (court_name in text2):
            hits.append(court_name)

    if not hits:
        m = re.search(r"(臺灣高等法院[一-龥]{1,4}分院)", text2)
        if m:
            hits.append(m.group(1))
        m = re.search(r"(臺灣[一-龥]{1,4}地方法院)", text2)
        if m:
            hits.append(m.group(1))
        m = re.search(r"(福建[一-龥]{1,4}地方法院)", text2)
        if m:
            hits.append(m.group(1))

    return _osc_unique_keep_order([_osc_normalize_court_name(x) for x in hits])


def _osc_normalize_case_word(s: str) -> str:
    word = str(s or "").strip()
    if not word:
        return ""
    word = word.replace("臺", "台")
    word = re.sub(r"\s+", "", word)
    word = re.sub(r"[字第號()（）【】\[\]{}「」『』《》,，.。:：;；/\\|_+\-]", "", word)
    return word


def _osc_parse_structured_case_spec(*, title: str = "", case_number: str = "") -> dict:
    texts = [
        ("case_number", str(case_number or "").strip()),
        ("title", str(title or "").strip()),
    ]
    out = {
        "case_year": "",
        "case_word": "",
        "case_no": "",
        "case_number_query": "",
        "case_marker": "",
        "courts": _osc_extract_court_names(title),
        "source": "",
    }
    patterns = [
        r"(?P<year>\d{2,3})\s*年度\s*(?P<word>[A-Za-z\u4e00-\u9fff]{1,16}?)\s*字?\s*第?\s*(?P<no>0*\d{1,8})\s*號?",
        r"(?P<year>\d{2,3})\s*(?P<word>[A-Za-z\u4e00-\u9fff]{1,16}?)\s*(?:字)?\s*第?\s*(?P<no>0*\d{1,8})\s*號?",
    ]
    for source, raw in texts:
        if not raw:
            continue
        text = raw.replace("（", "(").replace("）", ")")
        text = re.sub(r"\s+", "", text)
        for pat in patterns:
            m = re.search(pat, text)
            if not m:
                continue
            year = str(m.group("year") or "").strip()
            word = _osc_normalize_case_word(m.group("word") or "")
            no_raw = str(m.group("no") or "").strip()
            if (not year) or (not word) or (not no_raw):
                continue
            try:
                no = str(int(no_raw))
            except Exception:
                no = no_raw.lstrip("0") or "0"
            out.update(
                {
                    "case_year": year,
                    "case_word": word,
                    "case_no": no,
                    "case_number_query": f"{year}年度{word}字第{no}號",
                    "case_marker": f"{year}{word}{no}",
                    "source": source,
                }
            )
            return out
    return out


def _osc_extract_case_markers(s: str) -> set[str]:
    text = str(s or "")
    if not text:
        return set()
    pats = [
        r"\d{1,3}\s*[台臺]\s*[\u4e00-\u9fff]{1,10}\s*字?\s*第?\s*\d{1,8}\s*號?",
        r"\d{2,3}\s*年度\s*[\u4e00-\u9fff]{1,10}\s*字\s*第?\s*\d{1,8}\s*號?",
        r"\d{2,3}\s*年度\s*[\u4e00-\u9fff]{1,10}\s*第?\s*\d{1,8}\s*號?",
    ]
    out: set[str] = set()
    for p in pats:
        for m in re.findall(p, text):
            t = re.sub(r"\s+", "", m or "").replace("第", "").replace("號", "").replace("臺", "台")
            if t:
                out.add(t)
    return out


def _osc_load_judicial_search_results(result_obj: dict) -> list[dict]:
    items = []
    results_path = str(result_obj.get("results_path") or "").strip()
    if results_path and os.path.exists(results_path):
        try:
            with open(results_path, "r", encoding="utf-8") as f:
                cached = json.load(f) or {}
            items = cached.get("results") or []
        except Exception:
            items = []
    if not items:
        items = result_obj.get("results") or []
    return [it for it in items if isinstance(it, dict)]


def _osc_pick_exact_judicial_search_result(items: list[dict], *, title: str = "", case_number: str = "") -> dict:
    if not isinstance(items, list) or not items:
        return {}
    target = _osc_parse_structured_case_spec(title=title, case_number=case_number)
    if not (target.get("case_year") and target.get("case_word") and target.get("case_no")):
        return {}
    target_courts = {_osc_normalize_court_name(x) for x in (target.get("courts") or []) if str(x or "").strip()}
    target_title_norm = _osc_title_norm(title)
    target_word = _osc_normalize_case_word(target.get("case_word") or "")
    target_no = str(target.get("case_no") or "").lstrip("0") or "0"
    best = None
    best_score = -1.0
    for it in items:
        title2 = str(it.get("title") or "").strip()
        url2 = str(it.get("url") or "").strip()
        if not title2 or not url2:
            continue
        spec2 = _osc_parse_structured_case_spec(title=title2)
        if not (spec2.get("case_year") and spec2.get("case_word") and spec2.get("case_no")):
            continue
        if str(spec2.get("case_year") or "") != str(target.get("case_year") or ""):
            continue
        if _osc_normalize_case_word(spec2.get("case_word") or "") != target_word:
            continue
        no2 = str(spec2.get("case_no") or "").lstrip("0") or "0"
        if no2 != target_no:
            continue
        item_courts = {_osc_normalize_court_name(x) for x in _osc_extract_court_names(title2)}
        if target_courts and item_courts and target_courts.isdisjoint(item_courts):
            continue
        score = 10.0
        if target_courts and item_courts and (not target_courts.isdisjoint(item_courts)):
            score += 3.0
        if target_title_norm:
            title_norm2 = _osc_title_norm(title2)
            if title_norm2:
                ratio = difflib.SequenceMatcher(None, target_title_norm, title_norm2).ratio()
                score += ratio * 2.0
                if title_norm2 == target_title_norm:
                    score += 1.5
        if score > best_score:
            best_score = score
            best = it
    if not best:
        return {}
    out = dict(best)
    out["match_score"] = round(best_score, 4)
    return out


def _osc_fetch_fulltext_from_exact_case_search(*, title: str = "", case_number: str = "", timeout_sec: int = 180) -> dict:
    target = _osc_parse_structured_case_spec(title=title, case_number=case_number)
    if not (target.get("case_year") and target.get("case_word") and target.get("case_no")):
        return {"ok": False, "error": "structured_case_unavailable"}

    search_payload = {
        "keywords": "",
        "max_results": 20,
        "headless": True,
        "timeout_sec": max(60, min(240, int(timeout_sec))),
        "case_year": target.get("case_year") or "",
        "case_word": target.get("case_word") or "",
        "case_no": target.get("case_no") or "",
    }
    search_courts = [
        _OSC_JUDICIAL_COURT_SEARCH_LABELS.get(court_name, court_name)
        for court_name in (target.get("courts") or [])
        if str(court_name or "").strip()
    ]
    search_courts = _osc_unique_keep_order(search_courts)
    if search_courts:
        search_payload["courts"] = search_courts

    rr = _osc_run_skill(
        "judicial-web-search",
        _osc_skill_json_task("search", search_payload),
        timeout_sec=max(120, int(timeout_sec) + 60),
        route_key="osc:insights:fetch_full:exact_case_search",
    )
    rp = _osc_parse_skill_output(rr)
    if not rp.get("success"):
        return {"ok": False, "error": rp.get("error") or "judicial_exact_search_failed"}

    items = _osc_load_judicial_search_results(rp)
    best = _osc_pick_exact_judicial_search_result(items, title=title, case_number=case_number)
    if not best:
        return {"ok": False, "error": "judicial_exact_case_not_found"}

    fetch_payload = {
        "url": str(best.get("url") or "").strip(),
        "headless": True,
        "timeout_sec": max(45, min(180, int(timeout_sec))),
        "max_chars": 180000,
    }
    rr2 = _osc_run_skill(
        "judicial-web-search",
        _osc_skill_json_task("fetch_text", fetch_payload),
        timeout_sec=max(120, int(timeout_sec) + 60),
        route_key="osc:insights:fetch_full:exact_case_fetch",
    )
    rp2 = _osc_parse_skill_output(rr2)
    text = ""
    text_path = str(rp2.get("text_path") or "").strip() if isinstance(rp2, dict) else ""
    if text_path and os.path.exists(text_path):
        try:
            with open(text_path, "r", encoding="utf-8", errors="replace") as f:
                text = (f.read() or "").strip()
        except Exception:
            text = ""
    if (not text) and isinstance(rp2, dict):
        text = str(rp2.get("text") or "").strip()
    if (not text) and fetch_payload["url"]:
        direct = _osc_fetch_url_text(fetch_payload["url"], timeout=max(20, min(60, int(timeout_sec))))
        if direct.get("ok"):
            text = str(direct.get("text") or "").strip()
    if len(text) < 120:
        return {"ok": False, "error": "judicial_exact_case_text_not_found"}

    return {
        "ok": True,
        "source": "fallback_judicial_exact_case",
        "text": text,
        "matched": {
            "title": str(best.get("title") or ""),
            "url": str(best.get("url") or ""),
            "case_number_query": str(target.get("case_number_query") or ""),
        },
    }


def _osc_pick_best_manifest_item(items: list[dict], *, title: str = "", case_number: str = "") -> dict:
    if not isinstance(items, list) or not items:
        return {}
    target_title = str(title or "").strip()
    target_case = re.sub(r"\s+", "", str(case_number or ""))
    target_norm = _osc_title_norm(target_title)
    target_markers = _osc_extract_case_markers(target_title + " " + target_case)
    best = None
    best_score = -1.0
    for it in items:
        if not isinstance(it, dict):
            continue
        p = str(it.get("archived_text_path") or it.get("text_path") or "").strip()
        if not p or (not os.path.exists(p)):
            continue
        title2 = str(it.get("title") or "").strip()
        score = 0.0
        item_markers = _osc_extract_case_markers(title2)
        if target_case and (target_case in re.sub(r"\s+", "", title2)):
            score += 6.0
        if target_markers and item_markers:
            inter = target_markers.intersection(item_markers)
            if inter:
                score += 5.0
        n2 = _osc_title_norm(title2)
        if target_norm and n2:
            score += difflib.SequenceMatcher(None, target_norm, n2).ratio() * 2.5
        score += min(float(os.path.getsize(p) if os.path.exists(p) else 0) / 200000.0, 0.8)
        if score > best_score:
            best_score = score
            best = it
    if (not best) or best_score < 1.0:
        return {}
    out = dict(best)
    out["match_score"] = round(best_score, 4)
    return out


def _osc_summarize_legal_insight(full_text: str) -> str:
    text = str(full_text or "").strip()
    if not text:
        return ""
    from api.legal_research_quality import grounded_summary_evidence, validate_span_citations

    grounded = grounded_summary_evidence(text)
    evidence = grounded.get("evidence") or []
    if not grounded.get("ok") or not evidence:
        return "摘要失敗：裁判理由中沒有足夠的法院論理段落，未寫入實務見解庫。"

    evidence_block = "\n\n".join(
        f"[{span.get('span_id')}]\n{span.get('text')}"
        for span in evidence
        if isinstance(span, dict)
    )
    prompt = (
        "請把下列已從臺灣裁判全文精確擷取的法院理由，整理成結構化實務見解卡。\n"
        "只能使用所附 S1、S2…段落，不得補充模型記憶或杜撰裁判。\n"
        "不得把『原告主張／被告辯稱』當作法院見解。\n"
        "每個法律規則與法院涵攝句尾都必須標示支持段落，例如 [S2]。\n"
        "固定欄位：1) 法律爭點 2) 法律規則 3) 關鍵事實 4) 法院涵攝 "
        "5) 結果方向 6) 限制／例外 7) 核對狀態。\n"
        "沒有資料的欄位寫『原文未明示』，不要猜測。\n\n"
        f"{evidence_block}"
    )
    use_nim = str(os.environ.get("MAGI_LEGAL_SUMMARY_USE_NIM", "1")).strip().lower() not in {
        "0",
        "false",
        "no",
        "off",
    }
    if use_nim:
        try:
            from skills.bridge.nim_heavy import run_nim_chat

            result = run_nim_chat(
                prompt=prompt,
                task_type="legal_analysis",
                heavy=True,
                require_pii_scrub=True,
                data_classification="public_judgment",
                privacy_profile="public_judgment",
                restore_pii=False,
                timeout_sec=int(os.environ.get("OSC_INSIGHT_SUMMARY_TIMEOUT_SEC", "180") or "180"),
                system_prompt=(
                    "你是臺灣律師使用的裁判證據整理器。"
                    "所有命題必須由輸入的 S 編號原文支持；禁止臆測。"
                ),
            )
            out = str(result.get("response") or "").strip()
            validation = validate_span_citations(out, evidence)
            if (
                result.get("success")
                and validation.get("ok")
                and "法律爭點" in out
                and "法院涵攝" in out
                and "核對狀態" in out
            ):
                return out
        except Exception:
            logging.getLogger(__name__).debug(
                "grounded NIM legal summary unavailable",
                exc_info=True,
            )

    # Fail safe and lightweight: source-mapped extractive card is better than a
    # fluent but ungrounded local-model summary.
    return str(grounded.get("summary") or "")


def _osc_fetch_fulltext_from_judicial(*, title: str = "", case_number: str = "", case_reason: str = "", timeout_sec: int = 180) -> dict:
    """
    來源被登入保護/反爬阻擋時，
    先走司法院案號精準查詢，再退回司法院全文搜尋歸檔，
    最後才退回 judgment-collector。
    """
    reason = (title or "").strip() or (case_number or "").strip() or (case_reason or "").strip()
    if not reason:
        return {"ok": False, "error": "missing_query"}
    archive_payload = {
        "query": reason,
        "max_results": 5,
        "max_chars": 180000,
        "headless": True,
        "timeout_sec": max(90, min(240, int(timeout_sec))),
    }
    collect_payload = {
        "case_reason": reason,
        "case_number": (case_number or "").strip(),
        "max_results": 5,
        "max_chars": 180000,
        "headless": True,
        "timeout_sec": max(90, min(420, int(timeout_sec))),
        "save_to_db": True,
        "notify": False,
    }

    def _try_skill(skill: str, task: str, *, route_key: str, ok_source: str, summary_source: str, inline_source: str) -> dict:
        rr = _osc_run_skill(
            skill,
            task,
            timeout_sec=max(120, int(timeout_sec) + 60),
            route_key=route_key,
        )
        rp = _osc_parse_skill_output(rr)
        if not rp.get("success"):
            return {"ok": False, "error": rp.get("error") or f"{skill}_failed"}
        items = []
        manifest_path = str(rp.get("manifest_path") or "").strip()
        if manifest_path and os.path.exists(manifest_path):
            try:
                with open(manifest_path, "r", encoding="utf-8") as f:
                    mf = json.load(f) or {}
                items = mf.get("items") or []
            except Exception:
                items = []
        if not items:
            items = rp.get("items_preview") or rp.get("items") or []
        best = _osc_pick_best_manifest_item(items, title=title, case_number=case_number)
        path = str(best.get("archived_text_path") or best.get("text_path") or "").strip()
        if path and os.path.exists(path):
            try:
                with open(path, "r", encoding="utf-8", errors="replace") as f:
                    text = (f.read() or "").strip()
                if len(text) >= 120:
                    return {
                        "ok": True,
                        "source": ok_source,
                        "text": text,
                        "matched": {
                            "title": str(best.get("title") or ""),
                            "url": str(best.get("url") or ""),
                            "path": path,
                        },
                    }
            except Exception:
                logging.getLogger(__name__).debug("silent-catch at %s:%s", __name__, "_osc_fetch_fulltext_from_judicial:_try_skill:read", exc_info=True)
        summary_path = str(rp.get("summary_path") or "").strip()
        if summary_path and os.path.exists(summary_path):
            try:
                with open(summary_path, "r", encoding="utf-8", errors="replace") as f:
                    text = (f.read() or "").strip()
                if len(text) >= 120:
                    return {"ok": True, "source": summary_source, "text": text}
            except Exception:
                logging.getLogger(__name__).debug("silent-catch at %s:%s", __name__, "_osc_fetch_fulltext_from_judicial:_try_skill:summary", exc_info=True)
        for it in (items or []):
            txt = str(it.get("full_text") or it.get("text") or it.get("summary") or "").strip()
            if len(txt) >= 300:
                return {"ok": True, "source": inline_source, "text": txt}
        return {"ok": False, "error": f"{skill}_fulltext_not_found"}

    last_error = "judicial_fulltext_not_found"

    exact = _osc_fetch_fulltext_from_exact_case_search(
        title=title,
        case_number=case_number,
        timeout_sec=max(90, min(240, int(timeout_sec))),
    )
    if exact.get("ok"):
        return exact
    last_error = str(exact.get("error") or last_error)

    archive = _try_skill(
        "judicial-flow-search-archive",
        _osc_skill_json_task("search_archive", archive_payload),
        route_key="osc:insights:fetch_full:search_archive",
        ok_source="fallback_judicial_archive",
        summary_source="fallback_judicial_archive_summary",
        inline_source="fallback_judicial_archive_inline",
    )
    if archive.get("ok"):
        return archive
    last_error = str(archive.get("error") or last_error)

    collector = _try_skill(
        "judgment-collector",
        _osc_skill_json_task("collect", collect_payload),
        route_key="osc:insights:fetch_full:judgment_collector",
        ok_source="fallback_judgment_collector",
        summary_source="fallback_judgment_collector_summary",
        inline_source="fallback_judgment_collector_inline",
    )
    if collector.get("ok"):
        return collector
    last_error = str(collector.get("error") or last_error)
    return {"ok": False, "error": last_error}


# ---------------------------------------------------------------------------
# 8. Core helper functions
# ---------------------------------------------------------------------------

def _osc_norm_case_category(v: str) -> str:
    s = str(v or "").strip()
    if not s:
        return ""
    alias = {
        "法扶案件": "法律扶助案件",
        "法扶": "法律扶助案件",
        "法律扶助": "法律扶助案件",
    }
    return alias.get(s, s)


def _osc_resolve_case_id(ref: str) -> str:
    r = str(ref or "").strip()
    if not r:
        return ""
    row, _ = _osc_exec("SELECT id FROM cases WHERE id=%s LIMIT 1", (r,), fetch="one")
    if row and row.get("id"):
        return str(row.get("id"))
    row, _ = _osc_exec("SELECT id FROM cases WHERE case_number=%s LIMIT 1", (r,), fetch="one")
    if row and row.get("id"):
        return str(row.get("id"))
    return r


def _osc_safe_int(v, default=0) -> int:
    try:
        return int(v)
    except Exception:
        return int(default)


def _osc_truthy(v) -> bool:
    return str(v or "").strip().lower() in {"1", "true", "yes", "on", "y"}


def _osc_text(v):
    s = str(v if v is not None else "").strip()
    return s or None


def _osc_current_actor() -> str:
    try:
        if current_user.is_authenticated and getattr(current_user, "username", None):
            return str(current_user.username)
    except Exception:
        logging.getLogger(__name__).debug("silent-catch at %s:%s", __name__, "_osc_current_actor", exc_info=True)
    return "web"


def _osc_log_activity(action: str, entity_type: str = "", entity_id: str = "", details=None) -> None:
    try:
        payload = details
        if isinstance(payload, (dict, list)):
            payload = json.dumps(payload, ensure_ascii=False)
        payload = _osc_text(payload)
        _osc_exec(
            "INSERT INTO activity_logs (action, entity_type, entity_id, details, user) VALUES (%s,%s,%s,%s,%s)",
            (
                _osc_text(action) or "osc_action",
                _osc_text(entity_type),
                _osc_text(entity_id),
                payload,
                _osc_current_actor(),
            ),
            fetch="none",
        )
    except Exception as e:
        logger.warning("OSC activity log write failed: %s", e)


def _osc_accounting_window(today: Optional[date] = None) -> tuple[date, date]:
    today = today or date.today()
    if today.day <= 24:
        end_date = date(today.year, today.month, 24)
        if today.month == 1:
            start_date = date(today.year - 1, 12, 25)
        else:
            start_date = date(today.year, today.month - 1, 25)
    else:
        start_date = date(today.year, today.month, 25)
        if today.month == 12:
            end_date = date(today.year + 1, 1, 24)
        else:
            end_date = date(today.year, today.month + 1, 24)
    return start_date, end_date


def _osc_get_setting_value(setting_key: str, default: str = "") -> str:
    key = str(setting_key or "").strip()
    if not key:
        return str(default or "")
    try:
        row, _ = _osc_exec("SELECT value FROM settings WHERE `key`=%s", (key,), fetch="one")
        value = (row or {}).get("value")
        if value is None:
            return str(default or "")
        text = str(value).strip()
        return text if text else str(default or "")
    except Exception:
        return str(default or "")


def _osc_unique_strings(values) -> list[str]:
    seen = set()
    out = []
    for raw in values or []:
        s = str(raw or "").strip()
        if not s or s in seen:
            continue
        seen.add(s)
        out.append(s)
    return out


def _osc_read_plain_text(path: str) -> str:
    for enc in ("utf-8", "utf-8-sig", "cp950", "big5"):
        try:
            with open(path, "r", encoding=enc) as f:
                return f.read()
        except Exception:
            continue
    try:
        with open(path, "rb") as f:
            return f.read().decode("utf-8", errors="ignore")
    except Exception:
        return ""


def _osc_read_docx_text(path: str) -> str:
    try:
        from docx import Document  # type: ignore

        doc = Document(path)
        return "\n".join(p.text for p in doc.paragraphs if str(p.text or "").strip())
    except Exception:
        return ""


def _osc_read_pdf_text(path: str) -> str:
    try:
        from pypdf import PdfReader  # type: ignore

        reader = PdfReader(path)
        return "\n".join((page.extract_text() or "") for page in reader.pages)
    except Exception:
        logging.getLogger(__name__).debug("silent-catch at %s:%s", __name__, "_osc_read_pdf_text:pypdf", exc_info=True)
    try:
        from PyPDF2 import PdfReader  # type: ignore

        reader = PdfReader(path)
        return "\n".join((page.extract_text() or "") for page in reader.pages)
    except Exception:
        logging.getLogger(__name__).debug("silent-catch at %s:%s", __name__, "_osc_read_pdf_text:PyPDF2", exc_info=True)
    tool = shutil.which("pdftotext")
    if tool:
        try:
            proc = subprocess.run(
                [tool, "-layout", path, "-"],
                capture_output=True,
                text=True,
                encoding="utf-8",
                errors="ignore",
                check=False,
            )
            if proc.returncode == 0:
                return proc.stdout or ""
        except Exception:
            logging.getLogger(__name__).debug("silent-catch at %s:%s", __name__, "_osc_read_pdf_text:pdftotext", exc_info=True)
    return ""


def _osc_read_textutil_text(path: str) -> str:
    tool = shutil.which("textutil")
    if not tool:
        return ""
    try:
        proc = subprocess.run(
            [tool, "-convert", "txt", "-stdout", path],
            capture_output=True,
            text=True,
            encoding="utf-8",
            errors="ignore",
            check=False,
        )
        if proc.returncode == 0:
            return proc.stdout or ""
    except Exception:
        return ""
    return ""


def _osc_resolve_existing_local_path_with_candidates(raw_path: str) -> tuple[str, list[str]]:
    candidates = []
    norm = _osc_norm_path(raw_path).replace("\\", "/")
    if norm:
        candidates.append(norm)
    candidates.extend(_osc_local_path_candidates(raw_path))
    uniq = _osc_unique_strings(candidates)
    for cand in uniq:
        if cand and os.path.exists(cand):
            return cand, uniq
    return "", uniq


def _osc_read_reference_document(raw_path: str, max_chars: int = 9000) -> dict:
    resolved, candidates = _osc_resolve_existing_local_path_with_candidates(raw_path)
    meta = {
        "file_path": str(raw_path or "").strip(),
        "resolved_path": resolved,
        "candidates": candidates,
        "ok": False,
        "text": "",
        "error": "",
    }
    if not resolved:
        meta["error"] = "file_not_found"
        return meta

    ext = Path(resolved).suffix.lower()
    text = ""
    if ext == ".docx":
        text = _osc_read_docx_text(resolved) or _osc_read_textutil_text(resolved)
    elif ext == ".pdf":
        text = _osc_read_pdf_text(resolved)
    elif ext in {".txt", ".md", ".text", ".log", ".json", ".yaml", ".yml", ".csv", ".tsv", ".rst"}:
        text = _osc_read_plain_text(resolved)
    elif ext in {".doc", ".rtf", ".odt", ".html", ".htm"}:
        text = _osc_read_textutil_text(resolved)
    else:
        try:
            if os.path.getsize(resolved) <= 2_000_000:
                text = _osc_read_plain_text(resolved)
        except Exception:
            logging.getLogger(__name__).debug("silent-catch at %s:%s", __name__, "_osc_read_reference_document:fallback_read", exc_info=True)
        if not text:
            text = _osc_read_textutil_text(resolved)

    text = re.sub(r"\n{3,}", "\n\n", str(text or "")).strip()
    if len(text) > max_chars:
        text = text[:max_chars].rstrip() + "\n...[截斷]"
    meta["ok"] = bool(text)
    meta["text"] = text
    if not text:
        meta["error"] = "text_extract_failed"
    return meta
